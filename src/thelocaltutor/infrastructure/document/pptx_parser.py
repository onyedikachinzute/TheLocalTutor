"""PPTX text extractor using python-pptx."""

from __future__ import annotations

import logging

from thelocaltutor.core.exceptions import DocumentProcessingError
from thelocaltutor.infrastructure.document.base_parser import BaseParser, ParsedPage

log = logging.getLogger(__name__)


class PptxParser(BaseParser):
    def supports(self, file_path: str) -> bool:
        return file_path.lower().endswith(".pptx")

    def parse(self, file_path: str) -> list[ParsedPage]:
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError as exc:
            raise DocumentProcessingError("python-pptx is not installed. Run: pip install python-pptx") from exc

        pages: list[ParsedPage] = []
        try:
            prs = Presentation(file_path)
            for i, slide in enumerate(prs.slides, start=1):
                texts: list[str] = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            line = " ".join(run.text for run in para.runs).strip()
                            if line:
                                texts.append(line)
                    if hasattr(shape, "table"):
                        for row in shape.table.rows:
                            for cell in row.cells:
                                t = cell.text.strip()
                                if t:
                                    texts.append(t)
                if texts:
                    pages.append(ParsedPage(page_number=i, text="\n".join(texts)))
        except Exception as exc:
            raise DocumentProcessingError(f"Failed to parse PPTX: {exc}") from exc

        log.info("PPTX parsed: %d slides with text from %s", len(pages), file_path)
        return pages
